import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL

doc = docx.Document('task4_table.docx')

# 1. Fill Table 0 (Classification Table)
table_class = doc.tables[0]

# Add additional row if needed
while len(table_class.rows) < 6:
    table_class.add_row()

data_class = [
    ['Model', 'Accuracy', 'Precision', 'Recall', 'F1-score'],
    ['Logistic Regression', '72.65%', '75.09%', '68.56%', '71.68%'],
    ['Decision Tree', '72.44%', '76.24%', '65.95%', '70.72%'],
    ['Random Forest (Bagging)', '72.85%', '76.17%', '67.25%', '71.43%'],
    ['AdaBoost', '72.42%', '76.69%', '65.14%', '70.45%'],
    ['Gradient Boosting', '72.97%', '74.95%', '69.74%', '72.25%']
]

for r_idx, row_values in enumerate(data_class):
    row = table_class.rows[r_idx]
    for c_idx, val in enumerate(row_values):
        cell = row.cells[c_idx]
        cell.text = val
        # Format header row
        if r_idx == 0:
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.bold = True
                    run.font.size = Pt(10)
        else:
            for p in cell.paragraphs:
                if c_idx > 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.size = Pt(9.5)

# 2. Fill Table 1 (Regression Table - Mark as N/A since project is Classification)
table_reg = doc.tables[1]
data_reg = [
    ['Model', 'RSS', 'RMSE', 'R²'],
    ['N/A (Classification Project)', 'N/A', 'N/A', 'N/A'],
    ['N/A', 'N/A', 'N/A', 'N/A'],
    ['N/A', 'N/A', 'N/A', 'N/A']
]

for r_idx, row_values in enumerate(data_reg):
    row = table_reg.rows[r_idx]
    for c_idx, val in enumerate(row_values):
        cell = row.cells[c_idx]
        cell.text = val
        if r_idx == 0:
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.bold = True
                    run.font.size = Pt(10)
        else:
            for p in cell.paragraphs:
                if c_idx > 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.size = Pt(9.5)

# Save updated document
doc.save('task4_table.docx')
print("Successfully updated task4_table.docx with all model evaluation metrics!")
